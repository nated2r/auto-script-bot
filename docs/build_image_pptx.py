"""Build image-based 16:9 PPTX from generated slide PNGs."""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

SLIDE_DIR = Path(r"d:\各種好用的PY\自動產稿機器人\docs\ppt-slides")
OUT = Path(r"d:\各種好用的PY\自動產稿機器人\docs\自動產稿機器人_產品簡報_圖卡版.pptx")

FILES = [
    "slide-01-cover.png",
    "slide-02-pain.png",
    "slide-03-product.png",
    "slide-04-modes.png",
    "slide-05-rules.png",
    "slide-06-steps.png",
]


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for name in FILES:
        path = SLIDE_DIR / name
        if not path.exists():
            raise FileNotFoundError(path)
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(
            str(path),
            Inches(0),
            Inches(0),
            width=prs.slide_width,
            height=prs.slide_height,
        )
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
